#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Сборщик данных для сайта-конструктора.

Читает книгу ГМ (Manual_DM_RPG.xlsm) и превращает листы в data/data.js,
который подключает сайт. Учитываются не только значения ячеек, но и
цветовые легенды (заливка ячеек означает «раз в бой», «за единицу», «запретное» и т.п.)
и примечания к ячейкам (способности существ).

Запуск:
    pip install openpyxl cryptography
    python tools/build_data.py путь/к/Manual_DM_RPG.xlsm --password ПАРОЛЬ_ГМ

Результат:
    data/data.js   — открытые данные (расы, черты, дары, материалы, оружие, мир…);
    data/locked.js — закрытые разделы (бестиарий, навыки, социум, алхимия, артефакты,
                     механики), зашифрованные паролем ГМ (AES-256-GCM, ключ из пароля
                     через PBKDF2-SHA256). Без пароля их не прочитать даже из исходников.

Пароль можно не писать в командной строке: задайте переменную окружения GM_PASSWORD
или скрипт спросит его сам. Ключ --no-lock кладёт всё в открытый data.js.
"""
import base64
import datetime
import getpass
import gzip
import hashlib
import json
import math
import os
import re
import secrets
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

sys.path.insert(0, str(Path(__file__).resolve().parent))
from spelling import fix_text  # noqa: E402  исправление явных опечаток

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "data.js"
OUT_LOCKED = ROOT / "data" / "locked.js"
KDF_ITER = 250000

# ---------------------------------------------------------------- helpers

def num(v):
    """Число или None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else round(float(v), 4)
    s = str(v).strip().replace(",", ".")
    try:
        f = float(s)
        return int(f) if f.is_integer() else round(f, 4)
    except ValueError:
        return None


def txt(v):
    """Очищенный текст ячейки."""
    if v is None:
        return ""
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return ("%.4f" % v).rstrip("0").rstrip(".")
    s = str(v).replace("\r", "").strip()
    # Автор ставит ведущую точку/запятую, чтобы Excel не считал «+2» формулой
    s = re.sub(r"^[.,]\s*(?=[+\-−])", "", s)
    s = re.sub(r"[ \t]+", " ", s)
    s = re.sub(r"\n\s*", "\n", s)
    return fix_text(s)


def fill_key(cell):
    f = cell.fill
    if not f or not f.fill_type:
        return None
    c = f.fgColor
    if c is None:
        return None
    if c.type == "theme":
        return "th%d" % c.theme
    if c.type == "rgb":
        return c.rgb
    if c.type == "indexed":
        return "ix%d" % c.indexed
    return None


def col(c):
    return column_index_from_string(c) if isinstance(c, str) else c


def cells(ws, r, c1, c2):
    return [ws.cell(r, x) for x in range(col(c1), col(c2) + 1)]


def row_txt(ws, r, c1, c2):
    return [txt(c.value) for c in cells(ws, r, c1, c2)]


HEADER_FILL = "FFA5A5A5"
GREEN = "FFC6EFCE"
RED = "FFFFC7CE"

SIZE_NAMES = {
    "см": "Совсем малый", "м": "Малый", "с": "Средний", "н": "Нормальный",
    "б": "Большой", "г": "Гигантский", "т": "Титанический",
}


def strip_marks(name):
    """Название без служебных значков и список значков."""
    marks = [m for m in "♥♠♦♣◘" if m in name]
    clean = name
    for m in "♥♠♦♣◘":
        clean = clean.replace(m, "")
    return clean.strip(), marks


def generic_block(ws, r1, r2, c1, c2, header_row=None, skip_empty=True, keep_fill=None):
    """Прямоугольный блок → {columns, rows}. keep_fill: dict fill->tag для первой колонки."""
    c1, c2 = col(c1), col(c2)
    columns = []
    start = r1
    if header_row:
        columns = [txt(ws.cell(header_row, c).value) or "" for c in range(c1, c2 + 1)]
        start = header_row + 1
    rows = []
    for r in range(start, r2 + 1):
        vals = [txt(ws.cell(r, c).value) for c in range(c1, c2 + 1)]
        if skip_empty and not any(vals):
            continue
        row = {"c": vals}
        if keep_fill:
            fk = fill_key(ws.cell(r, c1))
            if fk in keep_fill:
                row["tag"] = keep_fill[fk]
        rows.append(row)
    # убрать пустые справа колонки
    width = max([len(columns)] + [max((i + 1 for i, v in enumerate(rw["c"]) if v), default=0) for rw in rows])
    columns = columns[:width]
    for rw in rows:
        rw["c"] = rw["c"][:width]
    return {"columns": columns, "rows": rows}


def sectioned(ws, r1, r2, c1, c2):
    """Лист «текстом»: строки с серой заливкой — заголовки разделов."""
    sections = []
    cur = None
    for r in range(r1, r2 + 1):
        vals = [txt(ws.cell(r, c).value) for c in range(col(c1), col(c2) + 1)]
        if not any(vals):
            continue
        if fill_key(ws.cell(r, col(c1))) == HEADER_FILL:
            cur = {"title": vals[0], "head": [v for v in vals[1:] if v], "rows": []}
            sections.append(cur)
            continue
        if cur is None:
            cur = {"title": "", "head": [], "rows": []}
            sections.append(cur)
        while vals and not vals[-1]:
            vals.pop()
        cur["rows"].append(vals)
    return sections


# ---------------------------------------------------------------- races

RACE_FREQ = {"th9": "раз в бой", "th4": "раз в день", "th6": "каждый ход/удар", "th5": "после убийства"}

SYMBOL_GENITIVE = {
    "огня": "огонь", "льда": "лёд", "воды": "вода", "земли": "земля", "песка": "песок",
    "природы": "природа", "пророды": "природа", "святости": "святость", "света": "свет",
    "тьмы": "тьма", "теней": "тьма", "крови": "кровь (2)", "молнии": "молния",
    "власти": "власть", "плоти": "плоть",
    "кристалла": "кристалл", "кристала": "кристалл", "иллюзии": "Иллюзия", "пепла": "пепел",
    "некро": "некро", "костей": "кость", "души": "душа (5)", "слизи": "слизь", "яда": "яд",
    "кислоты": "кислота", "плазмы": "плазма", "пси": "пси (3)", "истока": "исток (7)",
    "ткани": "ткань", "переноса": "перенос", "гравитации": "гравитация", "металла": "металл",
    "пара": "пар", "времени": "время", "лечения": "исцеление",
}

TALENT_ALIASES = {
    "шахтёрство": "шахтёрство", "фермерство": "фермерство", "ремесло": "ремесло",
    "кулинария": "кулинария", "свеживание": "свежевание", "свежевание": "свежевание",
    "скрытность": "скрытность", "воровство": "воровство", "внимательность": "внимательность",
    "точность": "точность", "харизма": "харизма", "харизмы": "харизма",
}

WEAPON_GROUPS = {
    "топоры": ["топор", "боевой топор"],
    "секиры": ["секира"], "секира": ["секира"],
    "рапиры": ["рапира"],
    "кинжалы": ["кинжал"],
    "копья": ["копьё"],
    "луки": ["пехотный лук", "лук", "длинный лук", "ужасающий лук"],
    "кастеты": ["кастет"],
    "катана": ["катана", "дай катана"], "катаны": ["катана", "дай катана"],
    "кхопешь": ["кхопеш"],
    "булава": ["булава"],
    "посохи": ["посох", "посох бо"],
    '"гвоздь"': ['"гвоздь"'],
    "макуакатль": ["макуауитль"],
    "арбалеты": ["наручный арбалет", "арбалет"],
}

COND_WORDS = ("в воде", "на суше", "на суще", "ночью", "днём", "днем", "в темноте", "на свету",
              "за союзника", "за каждого", "если ", "в тумане", "вне ", "в лабиринте", "на 2х", "на 2хода",
              "на 4х", "осёдлан", "в бою", "при ", "под ")


def parse_race_fx(text, is_penalty=False):
    """Числовые эффекты из текста бонуса/штрафа расы."""
    fx = []
    notes = []
    if not text:
        return fx
    low = text.lower().replace("ё", "ё")
    segs = [s.strip() for s in re.split(r"/", low) if s.strip()]
    for seg in segs:
        cond = any(w in seg for w in COND_WORDS)
        # харизма и таланты
        weapon_seg = any(g in seg for g in WEAPON_GROUPS)
        for m in re.finditer(r"(шахтёрство|фермерство|ремесло|кулинария|свеживание|свежевание|скрытность|воровство|внимательность|точность|харизм[аы]?)\s*([+\-−]\s*\d+)", seg):
            name = m.group(1)
            if name == "точность" and weapon_seg:
                continue
            if name.startswith("харизм"):
                name = "харизма"
            v = int(m.group(2).replace(" ", "").replace("−", "-"))
            fx.append({"k": "talent", "id": TALENT_ALIASES.get(name, name), "v": v, "cond": cond})
        # привлекательность
        for m in re.finditer(r"привлекательность\s*([+\-−]\s*\d+)", seg):
            fx.append({"k": "attr", "v": int(m.group(1).replace(" ", "").replace("−", "-")), "cond": cond})
        # зоркость
        if "ноч" not in seg:
            for m in re.finditer(r"зор[а-я]*\s*\+\s*(\d+)", seg):
                fx.append({"k": "zor", "v": int(m.group(1)), "cond": cond})
        # ОД
        # только «+2ОД», «ОД-1», «доп 3 ОД»; «через 9ОД», «за 1ОД» — длительность и цена, не бонус
        m = re.search(r"доп\.?\s*(\d+)\s*од\b|([+\-−]\s*\d+)\s*од\b|од\s*([+\-−]\s*\d+)", seg)
        if m:
            raw = (m.group(1) or m.group(2) or m.group(3)).replace(" ", "").replace("−", "-")
            fx.append({"k": "od", "v": int(raw), "cond": cond})
        # регенерации
        m = re.search(r"(?:хп реген|реген хп)\s*\+?\s*(\d+)", seg)
        if m:
            fx.append({"k": "hpRegen", "v": int(m.group(1)), "cond": cond})
        m = re.search(r"(?:мана реген|мп реген)\s*\+?\s*(\d+)\s*(х)?", seg)
        if m:
            if m.group(2):
                fx.append({"k": "mpRegenMul", "v": int(m.group(1)), "cond": cond})
            else:
                fx.append({"k": "mpRegen", "v": int(m.group(1)), "cond": cond})
        # оружие
        for g, types in WEAPON_GROUPS.items():
            m = re.search(re.escape(g) + r"\s*(?:точность\s*)?\+\s*(\d+)", seg)
            if m:
                kind = "weaponAcc" if "точност" in seg else "weapon"
                fx.append({"k": kind, "group": g.strip('"'), "types": types, "v": int(m.group(1)), "cond": cond})
        # символы
        for m in re.finditer(r"(?:символ|сивол|знак)\s+([а-яё]+)", seg):
            w = m.group(1)
            if w in ("по", "пренодлежности", "принодлежности", "бога"):
                fx.append({"k": "symbolByAffiliation"})
                continue
            if w in SYMBOL_GENITIVE:
                fx.append({"k": "symbol", "id": SYMBOL_GENITIVE[w]})
            elif w in ("дракона",):
                fx.append({"k": "symbol", "id": "дракон", "special": True})
            elif w in ("урабороса",):
                fx.append({"k": "symbol", "id": "Ураборос (символ)", "special": True})
        m = re.search(r"символ\s+слизи,\s*яда,\s*кислоты", seg)
        if m:
            fx += [{"k": "symbol", "id": "яд"}, {"k": "symbol", "id": "кислота"}]
        m = re.search(r"символы:\s*(.+)", seg)
        if m:
            for w in re.split(r",\s*", m.group(1)):
                w = w.strip()
                if w == "космос":
                    fx.append({"k": "symbol", "id": "Космос (6)"})
                elif w == "пси":
                    fx.append({"k": "symbol", "id": "пси (3)"})
                elif w:
                    fx.append({"k": "symbol", "id": w})
        # молитвы
        m = re.search(r"\+\s*(\d+)\s*к количеству молитв", seg)
        if m:
            fx.append({"k": "prayers", "v": int(m.group(1))})
        # доп. положительные черты
        m = re.search(r"могут взять ещё\s*(\d+)\s*положительн", seg)
        if m:
            fx.append({"k": "extraPos", "v": int(m.group(1))})
        # прочие пассивы (не таланты)
        for m in re.finditer(r"(алхимия|торговля|взлом|сбор трав|добыча|сбор)\s*\+?\s*(\d+)", seg):
            fx.append({"k": "passive", "id": m.group(1), "v": int(m.group(2))})
        if "язык зверей" in seg:
            fx.append({"k": "ability", "id": "язык зверей"})
        if "язык монстров" in seg:
            fx.append({"k": "ability", "id": "язык монстров"})
        if "шаманизм" in seg:
            fx.append({"k": "ability", "id": "Шаманизм"})
        if re.search(r"\b4 руки\b", seg):
            fx.append({"k": "hands", "v": 4})
        if "хобот как доп рука" in seg:
            fx.append({"k": "hands", "v": 3})
        if "не могут держать в руках оружия" in seg:
            fx.append({"k": "noWeapon"})
        if "не могут носить любые другие доспехи" in seg:
            fx.append({"k": "noArmor"})
        if "не имеют дебафа от роста" in seg:
            fx.append({"k": "noSizeDebuff"})
        if "средние доспехи не дают штраф" in seg:
            fx.append({"k": "noMediumArmorPenalty"})
        if "на начале не могут взять экипировку" in seg:
            fx.append({"k": "noStartGear"})
    # дедупликация символов
    seen = set()
    out = []
    for f in fx:
        key = json.dumps(f, ensure_ascii=False, sort_keys=True)
        if key in seen:
            continue
        seen.add(key)
        out.append(f)
    return out


# Ручные поправки к автоматическому разбору (проверено по тексту таблицы)
RACE_FX_FIX = {
    # «двойной ход» и «перемотка: +7ОД на 2 хода» — временные, не постоянные ОД
    "кселоры": {"drop_od": True},
    # «существа ниже уровнем теряют пси 15е,1ОД» — эффект на врагов
    "ми'го": {"drop_od": True},
    # «на суше -20л/-1ОД» — штраф только на суше
    "ангаймы": {"cond_pen_od": True},
    # у лунных эльфов «-30 ловкости» — способность (давление гравитации) на врагов
    # гадхарвы: «-15 всем статам» — только если долго не пьют
}


def round_level(x):
    """Уровни — целые: округление до ближайшего (0,5 — вверх)."""
    return int(math.floor(x + 0.5))


def build_races(wb):
    ws = wb["Расы"]
    races = []
    for r in range(2, ws.max_row + 1):
        name_raw = txt(ws.cell(r, 1).value)
        if not name_raw:
            continue
        name, marks = strip_marks(name_raw)
        stats = [num(ws.cell(r, c).value) or 0 for c in range(2, 7)]
        lvl = (sum(stats) - 25) / 3
        bonus = txt(ws.cell(r, 13).value)
        pen = txt(ws.cell(r, 14).value)
        race = {
            "id": name.lower(),
            "name": name,
            "marks": marks,
            "playable": not ("♥" in marks or "♠" in marks),
            "evolutionOnly": "♥" in marks,
            "npcOnly": "♠" in marks,
            "t": stats[0], "d": stats[1], "s": stats[2], "i": stats[3], "l": stats[4],
            "lvl": round_level(lvl),
            "align": txt(ws.cell(r, 8).value),
            "habitat": txt(ws.cell(r, 9).value),
            "faction": txt(ws.cell(r, 10).value),
            "religion": txt(ws.cell(r, 11).value),
            "size": txt(ws.cell(r, 12).value),
            "bonus": bonus,
            "penalty": pen,
            "evo": txt(ws.cell(r, 15).value),
            "bonusFreq": RACE_FREQ.get(fill_key(ws.cell(r, 13))),
            "fx": parse_race_fx(bonus),
            "pfx": parse_race_fx(pen, True),
        }
        fix = RACE_FX_FIX.get(race["id"])
        if fix and fix.get("drop_od"):
            race["fx"] = [f for f in race["fx"] if f["k"] != "od"]
        if fix and fix.get("cond_pen_od"):
            for f in race["pfx"]:
                if f["k"] == "od":
                    f["cond"] = True
        races.append(race)
    legend = {"♠": "не игровые расы", "♥": "расы, доступные только эволюцией"}
    return races, {"freq": list(RACE_FREQ.values()), "marks": legend}


# ---------------------------------------------------------------- society

JOIN = {"th6": "легко", "th9": "сложно", "th5": "скрыто", "th7": "нельзя вступить", "th4": "забытое"}
DEMONS = ["Diarkon", "Roara", "Xoria", "Grando", "Losis", "Horgen", "Gilltarro", "Estar"]


def build_society(wb):
    ws = wb["Социум"]
    groups = []
    for r in range(2, 72):
        name = txt(ws.cell(r, 1).value)
        if not name:
            continue
        g = {
            "row": r, "name": name,
            "st": row_txt(ws, r, "B", "F"),
            "status": txt(ws.cell(r, 7).value),
            "bonus": txt(ws.cell(r, 8).value),
            "type": txt(ws.cell(r, 9).value),
            "join": JOIN.get(fill_key(ws.cell(r, 1))),
        }
        g["demon"] = name.strip() in DEMONS
        groups.append(g)
    npcs = []
    for r in range(2, 72):
        nm = txt(ws.cell(r, col("L")).value)
        if not nm:
            continue
        npcs.append(row_txt(ws, r, "L", "T"))
    quests, gquests = [], []
    target = quests
    for r in range(2, 72):
        if txt(ws.cell(r, col("W")).value) == "Глобальные квесты" and not txt(ws.cell(r, col("V")).value):
            target = gquests
            continue
        v = row_txt(ws, r, "V", "Y")
        if any(v):
            target.append(v)
    note = txt(ws.cell(7, col("J")).value)
    return {
        "groups": groups, "npcs": npcs, "quests": quests, "globalQuests": gquests,
        "joinNote": note, "joinLegend": list(JOIN.values()),
    }


# ---------------------------------------------------------------- creatures

CREATURE_FILL = {"th9": "дракон", "th5": "босс", "th7": "легенда"}


def build_creatures(wb):
    ws = wb["Существа"]
    comments = {}
    for row in ws.iter_rows():
        for c in row:
            if c.comment:
                comments[(c.row, c.column)] = fix_text(c.comment.text.strip())
    heads = {}
    for r in range(1, ws.max_row + 1):
        if fill_key(ws.cell(r, 1)) == HEADER_FILL:
            heads[r] = txt(ws.cell(r, 1).value)
    section = "Основной список"
    out = []
    upgrades = []
    stat_names = ["тело", "дух", "сила", "интеллект", "ловкость"]
    for r in range(2, ws.max_row + 1):
        if r in heads:
            section = heads[r]
            continue
        if section == "улучшения монстров":
            v = row_txt(ws, r, "A", "G")
            if v[0]:
                upgrades.append(v)
            continue
        name_raw = txt(ws.cell(r, 1).value)
        if not name_raw:
            continue
        name, marks = strip_marks(name_raw)
        st = [num(ws.cell(r, c).value) for c in range(2, 7)]
        if all(x is None for x in st):
            continue
        st = [x or 0 for x in st]
        abil = []
        for ci in range(1, 15):
            t = comments.get((r, ci))
            if t:
                label = stat_names[ci - 2] if 2 <= ci <= 6 else "особое"
                abil.append([label, t])
        out.append({
            "n": name, "m": marks,
            "t": st[0], "d": st[1], "s": st[2], "i": st[3], "l": st[4],
            "lvl": round_level((sum(st) - 25) / 3),
            "ch": txt(ws.cell(r, 8).value),
            "hab": txt(ws.cell(r, 9).value),
            "sym": txt(ws.cell(r, 10).value),
            "sz": txt(ws.cell(r, 11).value),
            "drop": txt(ws.cell(r, 12).value),
            "note": txt(ws.cell(r, 13).value),
            "sec": section,
            "cat": CREATURE_FILL.get(fill_key(ws.cell(r, 1))),
            "ab": abil,
        })
    legend = [txt(ws.cell(r, 14).value) for r in range(2, 15) if txt(ws.cell(r, 14).value)]
    return out, upgrades, legend


# ---------------------------------------------------------------- skills sheet

SYMBOL_CAT = {"th5": "запретные знания", "th9": "порицаемая", "th4": "прокачивается", "th6": "изначально доступно"}


def build_skills(wb):
    ws = wb["скилы"]
    symbols = []
    for r in range(2, 70):
        name = txt(ws.cell(r, 1).value)
        if not name:
            continue
        v = row_txt(ws, r, "A", "G")
        symbols.append({
            "row": r, "name": name, "req": v[1], "e1": v[2], "e2": v[3], "link": v[4],
            "grim": v[5], "price": v[6],
            "cat": SYMBOL_CAT.get(fill_key(ws.cell(r, 1))) or "обычная",
        })
    utility = generic_block(ws, 75, 85, "A", "E")
    rituals = generic_block(ws, 87, 99, "A", "D", header_row=87)
    prayers = generic_block(ws, 101, 125, "A", "D", header_row=101)
    special = generic_block(ws, 129, 134, "A", "F")
    chords = generic_block(ws, 143, 151, "A", "D", header_row=143)
    stances = []
    for r in range(156, 179):
        v = row_txt(ws, r, "A", "D")
        if v[0]:
            stances.append({"name": v[0], "effect": v[1], "req": v[2], "price": v[3]})
    actives = []
    for r in range(185, 264):
        v = row_txt(ws, r, "A", "F")
        if v[0]:
            actives.append({"name": v[0], "effect": v[1], "req": v[2], "cd": v[3], "type": v[4], "price": v[5],
                            "special": fill_key(ws.cell(r, 1)) == "th5"})
    passives = []
    for r in range(270, 305):
        v = row_txt(ws, r, "A", "F")
        if v[0]:
            passives.append({"name": v[0], "effect": v[1], "req": v[2], "chance": v[3], "q": v[4], "price": v[5],
                             "base": fill_key(ws.cell(r, 1)) == "th6"})
    gifts = []
    for r in range(3, 50):
        v = row_txt(ws, r, "J", "L")
        if v[0]:
            gifts.append({"name": v[0], "bonus": v[1], "debuff": v[2]})
    curses_special = []
    for r in range(53, 59):
        v = row_txt(ws, r, "J", "K")
        if v[0]:
            curses_special.append({"name": v[0], "effect": v[1],
                                   "itemCurse": fill_key(ws.cell(r, col("J"))) == "th4"})
    curses = []
    for r in range(59, 97):
        v = row_txt(ws, r, "J", "K")
        if v[0]:
            curses.append({"name": v[0], "effect": v[1]})
    curse_rules = [txt(ws.cell(r, col("N")).value) for r in range(2, 9) if txt(ws.cell(r, col("N")).value)]
    buffs = []
    for r in range(99, 109):
        v = row_txt(ws, r, "J", "M")
        if v[0]:
            buffs.append({"name": v[0].replace("♥", "").strip(), "effect": v[1], "dur": v[2], "how": v[3],
                          "birth": "♥" in v[0]})
    mods = []
    for r in range(110, 152):
        v = row_txt(ws, r, "J", "M")
        if v[0]:
            mods.append(v)
    diseases = []
    DIS_CAT = {"th9": "расстройство (0 пси)", "th5": "великая болезнь", "th7": "особая"}
    for r in range(160, 210):
        v = row_txt(ws, r, "J", "L")
        if v[0]:
            diseases.append({"name": v[0].replace("♥", "").strip(), "e1": v[1], "e2": v[2],
                             "birth": "♥" in v[0], "cat": DIS_CAT.get(fill_key(ws.cell(r, col("J")))) or "обычная"})
    disease_note = txt(ws.cell(159, col("N")).value)
    traits_pos, traits_neg, traits_acq = [], [], []
    target = None
    for r in range(1, 56):
        a = txt(ws.cell(r, col("P")).value)
        b = txt(ws.cell(r, col("Q")).value)
        c = txt(ws.cell(r, col("R")).value)
        low = a.lower()
        if low == "положительные черты":
            target = traits_pos; continue
        if low == "негативные черты":
            target = traits_neg; continue
        if low == "приобретаймые":
            target = traits_acq; continue
        if not a or low.startswith("название"):
            continue
        if target is not None:
            item = {"name": a, "effect": b}
            if c:
                item["how"] = c
            target.append(item)
    legends = {
        "symbols": {"запретные знания": "запретная магия", "порицаемая": "порицаемая магия",
                    "изначально доступно": "изначально доступно", "прокачивается": "прокачивается",
                    "с рождения ♥": "доступно с рождения"},
        "curseRules": curse_rules,
    }
    return {
        "symbols": symbols, "utility": utility, "rituals": rituals, "prayers": prayers,
        "specialMagic": special, "chords": chords, "stances": stances, "actives": actives,
        "passives": passives, "gifts": gifts, "cursesSpecial": curses_special, "curses": curses,
        "buffs": buffs, "mods": mods, "diseases": diseases, "diseaseNote": disease_note,
        "traitsPos": traits_pos, "traitsNeg": traits_neg, "traitsAcq": traits_acq,
        "legends": legends,
    }


# ---------------------------------------------------------------- materials

SCALE = {"th9": "за предмет", "th6": "за единицу", "th4": "для каждого предмета отдельно",
         "th7": "не стакается / уникальный", "th5": "адское"}


def build_materials(wb):
    ws = wb["материаллы"]
    mats = []
    for r in range(2, 125):
        name = txt(ws.cell(r, 1).value)
        if not name:
            continue
        nm, marks = strip_marks(name)
        def eff(c):
            cell = ws.cell(r, c)
            return {"t": txt(cell.value), "sc": SCALE.get(fill_key(cell))}
        mats.append({
            "name": nm, "marks": marks,
            "dmg": num(ws.cell(r, 2).value), "arm": num(ws.cell(r, 3).value),
            "ew": eff(4), "ea": eff(5), "ej": eff(6),
            "src": txt(ws.cell(r, 7).value), "price": txt(ws.cell(r, 8).value),
            "skill": txt(ws.cell(r, 9).value),
        })
    crystals = []
    for r in range(129, 179):
        name = txt(ws.cell(r, 1).value)
        if not name:
            continue
        nm, marks = strip_marks(name)
        crystals.append({
            "name": nm, "marks": marks,
            "dmg": num(ws.cell(r, 2).value), "arm": num(ws.cell(r, 3).value),
            "ew": {"t": txt(ws.cell(r, 4).value)}, "ea": {"t": txt(ws.cell(r, 5).value)},
            "ej": {"t": txt(ws.cell(r, 6).value)},
            "src": txt(ws.cell(r, 7).value), "price": txt(ws.cell(r, 8).value),
            "skill": txt(ws.cell(r, 9).value),
        })
    notes = [txt(ws.cell(r, 10).value) for r in range(2, 10) if txt(ws.cell(r, 10).value)]
    return mats, crystals, {"scale": list(SCALE.values()), "notes": notes}


# ---------------------------------------------------------------- weapons & armor

def parse_req(s):
    """'2л,25с / 13' → ({'л':2,'с':25}, 13)."""
    s = (s or "").replace("ё", "е")
    parts = s.split("/")
    req = {}
    skill = None
    left = parts[0] if parts else ""
    for m in re.finditer(r"(\d+(?:[.,]\d+)?)\s*,?\s*([тдсилТДСИЛ])", left):
        req[m.group(2).lower()] = float(m.group(1).replace(",", "."))
    if len(parts) > 1:
        sk = re.search(r"\d+", parts[1])
        if sk:
            skill = int(sk.group(0))
    return req, skill


REQ_TOKEN = re.compile(r"(\d+(?:[.,]\d+)?)\s*,?\s*([тдсилТДСИЛ])")


def tidy_req(s):
    """'5,л,4с /8' → '5л, 4с / 8': единый вид требования; необычный текст остаётся как есть."""
    s = str(s or "").strip()
    parts = s.split("/")
    left = parts[0]
    toks = REQ_TOKEN.findall(left)
    if not toks or re.sub(r"[\s,.;]", "", REQ_TOKEN.sub("", left)):
        return s
    out = ", ".join(n.replace(".", ",") + l.lower() for n, l in toks)
    rest = [x.strip() for x in parts[1:]]
    return " / ".join([out] + rest) if rest else out


def build_gear(wb):
    ws = wb["оружие и броня"]
    weapons = []
    group = "ближний бой"
    for r in range(2, 96):
        name = txt(ws.cell(r, 1).value)
        if not name:
            continue
        if fill_key(ws.cell(r, 1)) == HEADER_FILL:
            group = {"стрелковое": "стрелковое", "снаряды": "снаряды"}.get(name, name)
            continue
        v = row_txt(ws, r, "A", "I")
        dmg_units = num(ws.cell(r, 2).value)
        special_units = None
        if dmg_units is None and "Х" in v[1].upper():
            m = re.match(r"([\d.,]+)\s*[XХхx]\s*(\d+)", v[1])
            if m:
                special_units = round(float(m.group(1).replace(",", ".")) * int(m.group(2)), 4)
                dmg_units = special_units
        req, skill = parse_req(v[3])
        g = group
        lname = name.lower()
        if group == "ближний бой":
            if lname in ("сфера", "гримуар", "жезл", "посох"):
                g = "магическое"
            elif lname in ("розарий", "кадильница", "молитвенник"):
                g = "молитвенное"
            elif lname in ("кирка", "лопата", "мотыга", "вилы", "капкан", "койло"):
                g = "инструмент"
        weapons.append({
            "name": name, "group": g,
            "u": dmg_units or 0, "un": num(ws.cell(r, 3).value) or 0, "uRaw": v[1],
            "reqRaw": tidy_req(v[3]), "req": req, "skill": skill,
            "fx": v[4], "dt": v[5], "range": v[6],
            "slots": num(ws.cell(r, 8).value), "hands": num(ws.cell(r, 9).value),
        })
    shields = []
    for r in range(97, 102):
        v = row_txt(ws, r, "A", "F")
        if v[0]:
            req, skill = parse_req(v[2])
            shields.append({"name": v[0], "u": num(ws.cell(r, 2).value), "uRaw": v[1], "reqRaw": tidy_req(v[2]),
                            "req": req, "skill": skill, "slots": num(ws.cell(r, 4).value),
                            "pen": num(ws.cell(r, 5).value) or 0, "fx": v[5],
                            "upgrade": v[1] == "улучшение"})
    armor = []
    for r in range(107, 131):
        v = row_txt(ws, r, "A", "E")
        if v[0]:
            req, skill = parse_req(v[2])
            nm = v[0].lower()
            cls = ("лёгкая" if nm.startswith("лёгк") else "средняя" if nm.startswith("средн") else
                   "тяжёлая" if nm.startswith("тяж") else "джаггернаут" if nm.startswith(("джагер", "джаггер")) else "одежда")
            armor.append({"name": v[0], "u": num(ws.cell(r, 2).value), "reqRaw": tidy_req(v[2]), "req": req,
                          "skill": skill, "slots": num(ws.cell(r, 4).value), "pen": num(ws.cell(r, 5).value) or 0,
                          "cls": cls})
    jewelry = generic_block(ws, 106, 114, "F", "I", header_row=106)
    jewelry_note = txt(ws.cell(105, col("F")).value)
    prosthetics = generic_block(ws, 131, 138, "A", "G", header_row=131)
    modern = generic_block(ws, 144, 152, "A", "B", header_row=144)
    modern_note = txt(ws.cell(144, col("E")).value)
    modules = generic_block(ws, 154, 161, "A", "C", header_row=154)
    instruments = generic_block(ws, 163, 178, "A", "D", header_row=163)
    gun_parts = generic_block(ws, 1, 33, "K", "O", header_row=1)
    gun_notes = [txt(ws.cell(r, col("R")).value) for r in range(1, 25) if txt(ws.cell(r, col("R")).value)]
    return {
        "weapons": weapons, "shields": shields, "armor": armor,
        "tables": {
            "jewelry": jewelry, "jewelryNote": jewelry_note, "prosthetics": prosthetics,
            "modern": modern, "modernNote": modern_note, "modules": modules, "instruments": instruments,
            "gunParts": gun_parts, "gunNotes": gun_notes,
        },
    }


# ---------------------------------------------------------------- alchemy

ALCH_NAME = {"th5": "запрещённое", "th7": "ценное", "th9": "наркотичное"}
ALCH_USE = {"th6": "огнестрел", "th4": "основа", "th7": "гримуар", "th9": "ковка/закалка", "th5": "трип-эффект",
            "FFF2F2F2": "крафт"}


def build_alchemy(wb):
    ws = wb["алхимия"]
    ingredients = []
    group = ""
    for r in range(2, 441):
        a = txt(ws.cell(r, 1).value)
        if not a:
            continue
        rest = [txt(ws.cell(r, c).value) for c in range(2, 8)]
        if fill_key(ws.cell(r, 1)) == HEADER_FILL or (not any(rest) and r in (2, 349, 382, 424)):
            group = a
            continue
        if not any(rest) and r > 2 and fill_key(ws.cell(r, 1)) is None and r in (2,):
            group = a
            continue
        effs = []
        for c in range(2, 6):
            cell = ws.cell(r, c)
            t = txt(cell.value)
            if t:
                fk = fill_key(cell)
                effs.append([t, "+" if fk == GREEN else "-" if fk == RED else ""])
        extra = []
        for c in range(8, 11):
            cell = ws.cell(r, c)
            t = txt(cell.value)
            if t:
                extra.append([t, ALCH_USE.get(fill_key(cell), "")])
        ingredients.append({
            "n": a, "tag": ALCH_NAME.get(fill_key(ws.cell(r, 1))), "g": group,
            "e": effs, "price": txt(ws.cell(r, 6).value), "lvl": txt(ws.cell(r, 7).value), "x": extra,
        })
    food = []
    for r in range(2, 38):
        v = row_txt(ws, r, "L", "Q")
        if v[0]:
            food.append({"name": v[0], "e": [x for x in v[1:4] if x], "cook": v[4], "price": v[5]})
    alcohol = []
    for r in range(49, 97):
        v = row_txt(ws, r, "L", "Q")
        if v[0]:
            alcohol.append({"name": v[0], "e": [x for x in v[1:4] if x], "cook": v[4], "price": v[5]})
    essences = []
    for r in range(2, 49):
        v = row_txt(ws, r, "S", "U")
        if v[0]:
            essences.append({"name": v[0], "effect": v[1], "price": v[2]})
    ess_note = txt(ws.cell(2, col("V")).value)
    return {
        "ingredients": ingredients, "food": food, "alcohol": alcohol, "essences": essences,
        "essenceNote": ess_note, "nameLegend": list(ALCH_NAME.values()), "useLegend": list(ALCH_USE.values()),
    }


# ---------------------------------------------------------------- misc sheet

def build_misc(wb):
    ws = wb["разное"]
    artifacts = []
    group = "артефакты"
    for r in range(2, 221):
        a = txt(ws.cell(r, 1).value)
        if not a:
            continue
        fk = fill_key(ws.cell(r, 1))
        if fk in ("th5", "th4") and not txt(ws.cell(r, 2).value):
            group = a.lower()
            continue
        v = row_txt(ws, r, "A", "E")
        artifacts.append({"name": v[0], "effect": v[1], "how": v[2], "from": v[3], "type": v[4], "g": group})
    items = []
    group = "предметы"
    for r in range(2, 72):
        a = txt(ws.cell(r, col("G")).value)
        if not a:
            continue
        fk = fill_key(ws.cell(r, col("G")))
        if fk in ("th7", "th8", "th9", "th6") and not txt(ws.cell(r, col("H")).value):
            group = a
            continue
        v = row_txt(ws, r, "G", "K")
        items.append({"name": v[0], "effect": v[1], "craft": v[2], "station": v[3], "price": v[4], "g": group})
    return artifacts, items


# ---------------------------------------------------------------- world & mechanics

def build_world(wb):
    ws = wb["миростроение"]
    layers = []
    for r in range(2, 22):
        code = txt(ws.cell(r, 1).value)
        if code:
            layers.append({"code": code, "name": txt(ws.cell(r, 2).value), "effect": txt(ws.cell(r, 3).value)})
    places = []
    region = ""
    for r in range(2, 20):
        a = txt(ws.cell(r, col("G")).value)
        if not a:
            continue
        if fill_key(ws.cell(r, col("G"))):
            region = a
        places.append({"name": a, "region": region, "monsters": txt(ws.cell(r, col("H")).value),
                       "players": txt(ws.cell(r, col("I")).value), "isRegion": bool(fill_key(ws.cell(r, col("G"))))})
    sections = sectioned(ws, 22, 168, "A", "E")
    return {"layers": layers, "places": places, "sections": sections}


def build_mechanics(wb):
    ws = wb["механики"]
    main = sectioned(ws, 1, 163, "A", "C")
    craft = generic_block(ws, 1, 28, "E", "G", header_row=1)
    dungeon = generic_block(ws, 1, 16, "I", "L", header_row=1)
    return {"main": main, "craft": craft, "dungeon": dungeon}


def build_schemes(wb):
    ws = wb["схемы"]
    grid = []
    for r in range(1, ws.max_row + 1):
        grid.append([txt(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)])
    while grid and not any(grid[-1]):
        grid.pop()
    while grid and not any(grid[0]):
        grid.pop(0)
    return grid


# ---------------------------------------------------------------- main

def encrypt_locked(payload, password):
    """AES-256-GCM, ключ = PBKDF2-SHA256(пароль). Формат совместим с Web Crypto в браузере."""
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    raw = gzip.compress(json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))
    salt = secrets.token_bytes(16)
    iv = secrets.token_bytes(12)
    key = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, KDF_ITER, 32)
    ct = AESGCM(key).encrypt(iv, raw, None)   # шифртекст + 16 байт тега
    b64 = lambda b: base64.b64encode(b).decode("ascii")
    return {"v": 1, "iter": KDF_ITER, "salt": b64(salt), "iv": b64(iv), "gz": True, "ct": b64(ct)}


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    flags = sys.argv[1:]
    if not args:
        print(__doc__)
        sys.exit(1)
    password = None
    if "--password" in flags:
        password = flags[flags.index("--password") + 1]
        args = [a for a in args if a != password]
    no_lock = "--no-lock" in flags
    if not no_lock and not password:
        password = os.environ.get("GM_PASSWORD") or getpass.getpass("Пароль ГМ для закрытых разделов: ")
    src = Path(args[0])
    wb = openpyxl.load_workbook(src, data_only=True)
    races, race_legend = build_races(wb)
    society = build_society(wb)
    creatures, upgrades, creature_legend = build_creatures(wb)
    skills = build_skills(wb)
    mats, crystals, mat_legend = build_materials(wb)
    gear = build_gear(wb)
    alchemy = build_alchemy(wb)
    artifacts, items = build_misc(wb)
    world = build_world(wb)
    mech = build_mechanics(wb)
    schemes = build_schemes(wb)

    # Всё, что нужно для создания персонажа, остаётся открытым — игрок собирает героя без ГМ.
    # Закрываются только справочные разделы целиком (дроп, обитание, цены, получение и т.п.).
    # 1) описания 11 талантов
    skills["talentInfo"] = [p for p in skills["passives"] if p.get("base")]
    # 2) «Предыдущий опыт»: физ. навыки и стойки (без цен)
    skills["startActives"] = [{k: a[k] for k in ("name", "effect", "req", "cd")} for a in skills["actives"]]
    skills["startStances"] = [{k: a[k] for k in ("name", "effect", "req")} for a in skills["stances"]]
    # 3) предки драконидов, таррэ и скордитов: существа без боссов и легенд, только статы и навыки
    insect_re = re.compile(r"инсекто|хитин|насеком|жук|оса|мух|муравь|богомол", re.I)
    ancestors = []
    for c in creatures:
        if c["sec"] in ("Боссы", "легенды"):
            continue
        a = {k: c[k] for k in ("n", "m", "t", "d", "s", "i", "l", "lvl", "ch", "sym", "sz", "sec", "ab")}
        if c["sec"] == "🕷":
            a["ins"] = "spider"
        elif insect_re.search((c.get("drop") or "") + " " + c["n"]):
            a["ins"] = "other"
        ancestors.append(a)
    # 4) демоны-предки тифлингов (символ принадлежности): только имя и описание из листа «Социум»
    demons = [{"name": g["name"].strip(), "bonus": g["bonus"]} for g in society["groups"] if g.get("demon")]
    # боги-покровители асимару: религии листа «Социум», кроме демонов и закрытых для вступления
    gods = [{"name": g["name"].strip(), "bonus": g["bonus"]} for g in society["groups"]
            if g["type"] == "религия" and not g.get("demon") and g.get("join") not in ("нельзя вступить", "забытое")]
    # 5) «Расхититель древностей»: случайный артефакт — обычный или проклятый (название и эффект)
    artifact_pool = [{"name": a["name"], "effect": a["effect"]} for a in artifacts if a["g"] == "артефакты"]
    artifact_pool += [{"name": a["name"], "effect": a["effect"], "cursed": True} for a in artifacts if a["g"] == "проклятые"]
    data = {
        "meta": {"source": src.name, "built": datetime.date.today().isoformat(), "locked": not no_lock},
        "races": races, "raceLegend": race_legend,
        "skills": skills,
        "materials": mats, "crystals": crystals, "materialLegend": mat_legend,
        "gear": gear,
        "alchemy": alchemy,
        "items": items,
        "ancestors": ancestors,
        "artifactPool": artifact_pool,
        "demons": demons,
        "gods": gods,
        "world": world, "schemes": schemes,
    }
    locked = {
        "creatures": creatures, "creatureUpgrades": upgrades, "creatureLegend": creature_legend,
        "society": society,
        "skills": {k: skills.pop(k) for k in ("stances", "actives", "passives")},
        "alchemy": {k: alchemy.pop(k) for k in ("ingredients", "alcohol", "essences", "essenceNote", "nameLegend", "useLegend")},
        "artifacts": artifacts,
        "mechanics": mech,
    }
    if no_lock:
        data.update({k: v for k, v in locked.items() if k not in ("skills", "alchemy")})
        data["skills"].update(locked["skills"])
        data["alchemy"].update(locked["alchemy"])
    js = "/* Сгенерировано tools/build_data.py из %s — не редактируйте вручную */\nwindow.GAME_DATA = %s;\n" % (
        src.name, json.dumps(data, ensure_ascii=False, separators=(",", ":")))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(js, encoding="utf-8")
    print("OK:", OUT, "%.0f KB" % (len(js.encode("utf-8")) / 1024))
    if not no_lock:
        enc = encrypt_locked(locked, password)
        ljs = "/* Закрытые разделы, зашифрованы паролем ГМ (tools/build_data.py) */\nwindow.GAME_LOCKED = %s;\n" % json.dumps(enc)
        OUT_LOCKED.write_text(ljs, encoding="utf-8")
        print("OK:", OUT_LOCKED, "%.0f KB" % (len(ljs) / 1024))
    elif OUT_LOCKED.exists():
        OUT_LOCKED.unlink()
    print("races", len(races), "creatures", len(creatures), "materials", len(mats), "crystals", len(crystals),
          "weapons", len(gear["weapons"]), "armor", len(gear["armor"]), "gifts", len(skills["gifts"]),
          "traits", len(skills["traitsPos"]), len(skills["traitsNeg"]), "artifacts", len(artifacts))


if __name__ == "__main__":
    main()
